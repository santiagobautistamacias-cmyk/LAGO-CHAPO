"""Figura de verificacion: los 7 perfiles extraidos, en orden hidraulico."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

wb = load_workbook("/projects/sandbox/LAGO-CHAPO/Perfiles_RioNegro_analisis.xlsx")

# orden hidraulico y thalweg desde la hoja Progresivas
orden, thal = [], {}
ws = wb["Progresivas"]
for row in ws.iter_rows(min_row=5, max_row=11, max_col=3, values_only=True):
    if row[1]:
        orden.append(str(row[1]))
        thal[str(row[1])] = row[2]

perfiles = {}
ws = wb["Perfiles"]
for plano, _, corte, _, est, cota in ws.iter_rows(min_row=2, values_only=True):
    if plano != "2018-11":
        continue
    perfiles.setdefault(str(corte), []).append((est, cota))

fig, axes = plt.subplots(4, 2, figsize=(15, 15))
axes = axes.ravel()

for ax, nom in zip(axes, orden):
    p = sorted(perfiles[nom])
    x = [q[0] for q in p]
    z = [q[1] for q in p]
    ax.plot(x, z, "-o", ms=3, lw=1.8, color="#1F4E79")
    ax.fill_between(x, min(z) - 1, z, color="#1F4E79", alpha=0.10)
    for niv, col, et in [(230, "#C00000", "230 (mín. histórico)"),
                         (231, "#ED7D31", "231 (mín. 3TA)"),
                         (243, "#2E75B6", "243 (máx. normal)")]:
        if min(z) - 1 <= niv <= max(z) + 1:
            ax.axhline(niv, color=col, ls="--", lw=1.1, alpha=0.85)
            ax.text(x[-1], niv, f" {et}", va="center", fontsize=7, color=col)
    ax.set_title(f"CORTE {nom}   —   thalweg {thal[nom]:.2f} m.s.n.m.",
                 fontsize=11, fontweight="bold")
    ax.set_xlabel("Estación (m)", fontsize=9)
    ax.set_ylabel("Cota (m.s.n.m.)", fontsize=9)
    ax.grid(alpha=0.3)
    ax.margins(x=0.02)

axes[-1].axis("off")
axes[-1].text(0.02, 0.95,
              "Río Negro — Lago Chapo\n"
              "Perfiles extraídos de los DWG de monitoreo\n\n"
              "Orden: aguas abajo → aguas arriba\n"
              f"({' → '.join(orden)})\n\n"
              "Escala corregida por exageración vertical ×10.\n"
              "Geometría idéntica en los planos de 2014 y 2018.\n\n"
              "Solo los cortes 1 y A tienen el lecho bajo la\n"
              "cota 230, es decir en la zona de influencia del\n"
              "abatimiento del lago, que es donde el informe\n"
              "de INGETEC sitúa la erosión retrogradante.",
              va="top", fontsize=10, family="monospace")

fig.suptitle("Secciones transversales de monitoreo — Río Negro, Lago Chapo",
             fontsize=14, fontweight="bold", y=0.997)
fig.tight_layout(rect=[0, 0, 1, 0.985])
salida = "/projects/sandbox/LAGO-CHAPO/perfiles_rio_negro.png"
fig.savefig(salida, dpi=140)
print("->", salida)
